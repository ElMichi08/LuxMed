from __future__ import annotations
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
from datetime import datetime, date
from app.domain.entities import Paciente, EstadoValidacion
from app.domain.ports import IExcelHandler

class ExcelHandler(IExcelHandler):
    def __init__(self) -> None:
        self._nombre_hoja = "BASE"
        self._relleno_rojo = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")

    def leer_pacientes(self, ruta_archivo: str) -> list[Paciente]:
        pacientes = []
        df = pd.read_excel(ruta_archivo, sheet_name=self._nombre_hoja, dtype=str)
        
        for _, fila in df.iterrows():
            nombre = str(fila.get("NOMBRE Y APELLIDOS", "")).strip()
            cedula = str(fila.get("CEDULA", "")).strip()
            aporta = str(fila.get("APORTA", "")).strip()
            establecimiento = str(fila.get("NOM ESTABLECIMIENTO", "")).strip()
            
            f_nac_raw = str(fila.get("FECHA NACIMIENTO", "")).strip()
            f_at_raw = str(fila.get("FECHA ATENCION", "")).strip()
            
            try:
                dt_nac = _parsear_fecha(f_nac_raw)
                fecha_nacimiento = dt_nac.date()
            except (ValueError, TypeError):
                fecha_nacimiento = date(1900, 1, 1)
                aporta = "FECHA ERRÓNEA"

            try:
                dt_at = _parsear_fecha(f_at_raw)
                fecha_atencion = dt_at.date()
            except (ValueError, TypeError):
                fecha_atencion = date.today()

            pacientes.append(Paciente(
                nombre_y_apellidos=nombre,
                cedula=cedula,
                fecha_nacimiento=fecha_nacimiento,
                aporta=aporta,
                fecha_atencion=fecha_atencion,
                nom_establecimiento=establecimiento
            ))
        return pacientes

    def exportar_excel_limpio(self, ruta_destino: str, pacientes: list[Paciente]) -> None:
        datos = []
        validos = [p for p in pacientes if p.estado == EstadoValidacion.VALIDO]
        
        for p in validos:
            datos.append({
                "NOMBRE Y APELLIDOS": p.nombre_y_apellidos,
                "CEDULA": p.cedula,
                "FECHA NACIMIENTO": p.fecha_nacimiento.strftime("%d/%m/%Y"),
                "APORTA": p.aporta,
                "FECHA ATENCION": p.fecha_atencion.strftime("%d/%m/%Y"),
                "NOM ESTABLECIMIENTO": p.nom_establecimiento
            })
            
        df_salida = pd.DataFrame(datos)
        df_salida.to_excel(ruta_destino, sheet_name=self._nombre_hoja, index=False)

    def exportar_excel_auditoria(self, ruta_origen: str, ruta_destino: str, pacientes: list[Paciente]) -> None:
        wb = openpyxl.load_workbook(ruta_origen)
        if self._nombre_hoja not in wb.sheetnames:
            return
            
        ws = wb[self._nombre_hoja]
        mapeo_pacientes = {p.cedula: p for p in pacientes}
        
        headers = [str(cell.value).strip() for cell in ws[1]]
        try:
            col_cedula = headers.index("CEDULA") + 1
            col_aporta = headers.index("APORTA") + 1
        except ValueError:
            wb.save(ruta_destino)
            return

        for row_idx in range(2, ws.max_row + 1):
            cedula_celda = str(ws.cell(row=row_idx, column=col_cedula).value).strip()
            
            if cedula_celda in mapeo_pacientes:
                paciente = mapeo_pacientes[cedula_celda]
                
                ws.cell(row=row_idx, column=col_aporta, value=paciente.aporta)
                
                if paciente.es_auditoria_rojo:
                    for col_idx in range(1, ws.max_column + 1):
                        ws.cell(row=row_idx, column=col_idx).fill = self._relleno_rojo
                        
        wb.save(ruta_destino)


FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d")


def _parsear_fecha(texto: str) -> datetime:
    for formato in FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    raise ValueError(texto)
